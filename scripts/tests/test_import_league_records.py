import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from scripts.import_league_records import (
    assert_no_contact_values,
    import_records,
)


class ImportLeagueRecordsTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.workspace = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_import_extracts_members_and_removes_registration_sheet(self) -> None:
        source = self.workspace / "records.xlsx"
        sanitized = self.workspace / "sanitized.xlsx"
        private = self.workspace / "league-members.json"
        self._make_source_workbook(source)

        summary = import_records(source, sanitized, private)

        self.assertEqual(summary.member_count, 1)
        self.assertEqual(summary.removed_sheet, "2022")
        self.assertEqual(summary.retained_sheets, ("Winners", "2025"))
        self.assertEqual(load_workbook(sanitized).sheetnames, ["Winners", "2025"])
        payload = json.loads(private.read_text(encoding="utf-8"))
        self.assertEqual(payload["source"], "records.xlsx")
        self.assertEqual(payload["sourceSheet"], "2022")
        self.assertEqual(
            payload["members"],
            [
                {
                    "sourceRow": 2,
                    "firstName": "Test",
                    "lastName": "Member",
                    "email": "test@example.com",
                    "phone": "5550100000",
                    "teamNameReservation": "Test Team",
                    "duesPaid": "Yes",
                }
            ],
        )

    def test_import_rejects_workbook_without_registration_sheet(self) -> None:
        source = self.workspace / "records.xlsx"
        sanitized = self.workspace / "sanitized.xlsx"
        private = self.workspace / "league-members.json"
        workbook = Workbook()
        workbook.active.title = "Winners"
        workbook.save(source)

        with self.assertRaisesRegex(ValueError, "2022 registration sheet"):
            import_records(source, sanitized, private)

        self.assertFalse(sanitized.exists())
        self.assertFalse(private.exists())

    def test_contact_scan_rejects_email_in_retained_sheet(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Winners"
        worksheet["A1"] = "winner@example.com"

        with self.assertRaisesRegex(ValueError, "contact-like value"):
            assert_no_contact_values(workbook)

    @staticmethod
    def _make_source_workbook(path: Path) -> None:
        workbook = Workbook()
        winners = workbook.active
        winners.title = "Winners"
        winners.append(["Champion", "Season"])
        winners.append(["Example Champion", 2025])

        registration = workbook.create_sheet("2022")
        registration.append(
            [
                None,
                "First",
                "Last",
                "Email",
                "Phone # (Optional - for trades and smack talk)",
                "Love Language",
                "Team Name Reservation",
                "Dues paid",
                "send to benmo69420",
                "Draft dates that work for you",
            ]
        )
        registration.append(
            [
                None,
                "Test",
                "Member",
                "test@example.com",
                5_550_100_000,
                None,
                "Test Team",
                "Yes",
                None,
                None,
            ]
        )
        registration.append([None, None, None, None, None, None, None, None])

        current = workbook.create_sheet("2025")
        current.append(["Place", "Team"])
        current.append([1, "Safe Team"])
        workbook.save(path)


if __name__ == "__main__":
    unittest.main()
