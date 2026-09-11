#!/usr/bin/env python3
"""Extract private league-member data and sanitize the history workbook."""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


REGISTRATION_SHEET = "2022"
EMAIL_PATTERN = re.compile(r"[^\s@]+@[^\s@]+\.[^\s@]+")
MEMBER_COLUMNS = {
    "firstName": 2,
    "lastName": 3,
    "email": 4,
    "phone": 5,
    "loveLanguage": 6,
    "teamNameReservation": 7,
    "duesPaid": 8,
    "draftAvailability": 10,
}


@dataclass(frozen=True)
class ImportSummary:
    member_count: int
    removed_sheet: str
    retained_sheets: tuple[str, ...]


def clean_value(value: Any, *, phone: bool = False) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
    if isinstance(value, (date, datetime)):
        value = value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        value = int(value)
    if phone:
        return str(value)
    return value


def extract_members(worksheet: Worksheet) -> list[dict[str, Any]]:
    members: list[dict[str, Any]] = []
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=2, max_col=max(MEMBER_COLUMNS.values())),
        start=2,
    ):
        first_name = clean_value(row[MEMBER_COLUMNS["firstName"] - 1].value)
        last_name = clean_value(row[MEMBER_COLUMNS["lastName"] - 1].value)
        if first_name is None and last_name is None:
            continue

        member: dict[str, Any] = {"sourceRow": row_number}
        for field, column in MEMBER_COLUMNS.items():
            value = clean_value(row[column - 1].value, phone=field == "phone")
            if value is not None:
                member[field] = value
        members.append(member)
    return members


def is_contact_like(value: Any) -> bool:
    if value is None:
        return False
    text = str(value).strip()
    if EMAIL_PATTERN.search(text):
        return True
    digits = re.sub(r"\D", "", text)
    return len(digits) in (10, 11)


def assert_no_contact_values(workbook: Workbook) -> None:
    leaks: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if is_contact_like(cell.value):
                    leaks.append(f"{worksheet.title}!{cell.coordinate}")
    if leaks:
        locations = ", ".join(leaks[:10])
        raise ValueError(f"Found contact-like value in retained history: {locations}")


def write_private_directory(
    output: Path,
    source_name: str,
    members: list[dict[str, Any]],
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "source": source_name,
        "sourceSheet": REGISTRATION_SHEET,
        "members": members,
    }
    output.write_text(
        json.dumps(payload, indent=2, ensure_ascii=True) + "\n",
        encoding="utf-8",
    )


def import_records(
    source: Path,
    sanitized_output: Path,
    private_output: Path,
) -> ImportSummary:
    workbook = load_workbook(source)
    if REGISTRATION_SHEET not in workbook.sheetnames:
        raise ValueError(
            f"Workbook does not contain the {REGISTRATION_SHEET} registration sheet"
        )

    registration = workbook[REGISTRATION_SHEET]
    members = extract_members(registration)
    workbook.remove(registration)
    assert_no_contact_values(workbook)

    sanitized_output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(sanitized_output)
    write_private_directory(private_output, source.name, members)

    return ImportSummary(
        member_count=len(members),
        removed_sheet=REGISTRATION_SHEET,
        retained_sheets=tuple(workbook.sheetnames),
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path, help="Source or sanitized XLSX workbook")
    parser.add_argument("--sanitized-output", type=Path)
    parser.add_argument("--private-output", type=Path)
    parser.add_argument(
        "--check",
        action="store_true",
        help="Check a workbook for contact-like values without writing files",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.check:
        workbook = load_workbook(args.source, read_only=True, data_only=True)
        assert_no_contact_values(workbook)
        print(json.dumps({"workbook": str(args.source), "contactLikeValues": 0}))
        return

    if args.sanitized_output is None or args.private_output is None:
        raise SystemExit(
            "--sanitized-output and --private-output are required for import"
        )

    summary = import_records(args.source, args.sanitized_output, args.private_output)
    print(
        json.dumps(
            {
                "memberCount": summary.member_count,
                "removedSheet": summary.removed_sheet,
                "retainedSheets": summary.retained_sheets,
            }
        )
    )


if __name__ == "__main__":
    main()
