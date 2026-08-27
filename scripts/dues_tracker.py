#!/usr/bin/env python3
"""Local-only 2026 league dues tracker."""

from __future__ import annotations

import argparse
import json
import re
import webbrowser
from datetime import datetime
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlparse

from openpyxl import load_workbook


EDITABLE_FIELDS = {"paid", "venmoHandle", "notes"}
DUES_EXEMPTIONS = {"ben-ray": "Commissioner"}


def member_key(first_name: str, last_name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", f"{first_name}{last_name}".lower())


def member_id(first_name: str, last_name: str) -> str:
    words = re.findall(r"[a-z0-9]+", f"{first_name} {last_name}".lower())
    return "-".join(words)


def seed_members(history_path: Path, contacts_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(history_path, read_only=True, data_only=True)
    if "2025" not in workbook.sheetnames:
        raise ValueError("History workbook does not contain the 2025 roster sheet")

    contact_payload = json.loads(contacts_path.read_text(encoding="utf-8"))
    contacts = {
        member_key(member.get("firstName", ""), member.get("lastName", "")): member
        for member in contact_payload.get("members", [])
    }

    members: list[dict[str, Any]] = []
    worksheet = workbook["2025"]
    for row_number in range(2, 1000):
        first_value = worksheet.cell(row_number, 3).value
        last_value = worksheet.cell(row_number, 4).value
        if first_value is None or last_value is None:
            continue
        first_name = str(first_value).strip()
        last_name = str(last_value).strip()
        contact = contacts.get(member_key(first_name, last_name), {})
        current_member_id = member_id(first_name, last_name)
        exemption_reason = DUES_EXEMPTIONS.get(current_member_id)
        members.append(
            {
                "id": current_member_id,
                "name": f"{first_name} {last_name}",
                "firstName": first_name,
                "lastName": last_name,
                "email": contact.get("email"),
                "phone": contact.get("phone"),
                "venmoHandle": "",
                "paid": False,
                "paidAt": None,
                "notes": "",
                "duesRequired": exemption_reason is None,
                "exemptionReason": exemption_reason,
            }
        )
    return members


def write_tracker(path: Path, tracker: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = path.with_suffix(path.suffix + ".tmp")
    temporary_path.write_text(
        json.dumps(tracker, indent=2, ensure_ascii=True) + "\n",
        encoding="utf-8",
    )
    temporary_path.replace(path)


def load_tracker(
    history_path: Path,
    contacts_path: Path,
    state_path: Path,
    *,
    season: int,
) -> dict[str, Any]:
    existing_members: dict[str, dict[str, Any]] = {}
    if state_path.exists():
        existing = json.loads(state_path.read_text(encoding="utf-8"))
        if existing.get("season") == season:
            existing_members = {
                member["id"]: member for member in existing.get("members", [])
            }

    members = seed_members(history_path, contacts_path)
    for member in members:
        previous = existing_members.get(member["id"], {})
        preserved_fields = ["venmoHandle", "notes"]
        if member["duesRequired"]:
            preserved_fields.extend(["paid", "paidAt"])
        for field in preserved_fields:
            if field in previous:
                member[field] = previous[field]

    tracker = {"season": season, "members": members}
    write_tracker(state_path, tracker)
    return tracker


def update_member(
    state_path: Path,
    target_member_id: str,
    changes: dict[str, Any],
    *,
    now: str | None = None,
) -> dict[str, Any]:
    tracker = json.loads(state_path.read_text(encoding="utf-8"))
    member = next(
        (item for item in tracker["members"] if item["id"] == target_member_id),
        None,
    )
    if member is None:
        raise KeyError(target_member_id)

    unknown_fields = set(changes) - EDITABLE_FIELDS
    if unknown_fields:
        raise ValueError(f"Unsupported fields: {', '.join(sorted(unknown_fields))}")

    if "paid" in changes:
        paid = changes["paid"]
        if not isinstance(paid, bool):
            raise ValueError("paid must be a boolean")
        if paid and not member.get("duesRequired", True):
            raise ValueError("exempt member cannot be marked paid")
        member["paid"] = paid
        member["paidAt"] = (
            now or datetime.now().astimezone().isoformat(timespec="seconds")
            if paid
            else None
        )

    if "venmoHandle" in changes:
        handle = changes["venmoHandle"]
        if not isinstance(handle, str):
            raise ValueError("venmoHandle must be a string")
        member["venmoHandle"] = handle.strip().lstrip("@").strip()

    if "notes" in changes:
        notes = changes["notes"]
        if not isinstance(notes, str):
            raise ValueError("notes must be a string")
        member["notes"] = notes.strip()

    write_tracker(state_path, tracker)
    return member


def create_server(
    history_path: Path,
    contacts_path: Path,
    state_path: Path,
    *,
    season: int,
    port: int,
) -> ThreadingHTTPServer:
    load_tracker(history_path, contacts_path, state_path, season=season)

    class DuesRequestHandler(BaseHTTPRequestHandler):
        def send_bytes(
            self,
            body: bytes,
            content_type: str,
            status: HTTPStatus = HTTPStatus.OK,
        ) -> None:
            self.send_response(status)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header(
                "Content-Security-Policy",
                "default-src 'self'; style-src 'self' 'unsafe-inline'; "
                "script-src 'self' 'unsafe-inline'; connect-src 'self'; "
                "img-src 'self' data:; base-uri 'none'; frame-ancestors 'none'",
            )
            self.end_headers()
            self.wfile.write(body)

        def send_json(self, payload: Any, status: HTTPStatus = HTTPStatus.OK) -> None:
            body = json.dumps(payload).encode("utf-8")
            self.send_bytes(body, "application/json; charset=utf-8", status)

        def do_GET(self) -> None:
            path = urlparse(self.path).path
            if path == "/":
                html_path = Path(__file__).with_name("dues_tracker.html")
                self.send_bytes(html_path.read_bytes(), "text/html; charset=utf-8")
                return
            if path == "/api/dues":
                tracker = json.loads(state_path.read_text(encoding="utf-8"))
                self.send_json(tracker)
                return
            self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)

        def do_PATCH(self) -> None:
            path = urlparse(self.path).path
            prefix = "/api/members/"
            if not path.startswith(prefix):
                self.send_json({"error": "Not found"}, HTTPStatus.NOT_FOUND)
                return

            try:
                content_length = int(self.headers.get("Content-Length", "0"))
                if content_length > 65_536:
                    raise ValueError("Request body is too large")
                changes = json.loads(self.rfile.read(content_length) or b"{}")
                if not isinstance(changes, dict):
                    raise ValueError("Request body must be a JSON object")
                updated = update_member(
                    state_path,
                    unquote(path[len(prefix) :]),
                    changes,
                )
            except KeyError as error:
                self.send_json(
                    {"error": f"Unknown member: {error.args[0]}"},
                    HTTPStatus.NOT_FOUND,
                )
                return
            except (ValueError, json.JSONDecodeError) as error:
                self.send_json({"error": str(error)}, HTTPStatus.BAD_REQUEST)
                return

            self.send_json(updated)

        def log_message(self, format: str, *args: Any) -> None:
            return

    return ThreadingHTTPServer(("127.0.0.1", port), DuesRequestHandler)


def parse_args() -> argparse.Namespace:
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--season", type=int, default=2026)
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument(
        "--history",
        type=Path,
        default=root / "history/league/ultimate-guillotine-records.xlsx",
    )
    parser.add_argument(
        "--contacts",
        type=Path,
        default=root / "data/private/league-members.json",
    )
    parser.add_argument(
        "--state",
        type=Path,
        default=root / "data/private/dues-2026.json",
    )
    parser.add_argument("--no-open", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    server = create_server(
        args.history,
        args.contacts,
        args.state,
        season=args.season,
        port=args.port,
    )
    url = f"http://127.0.0.1:{server.server_port}"
    print(f"{args.season} dues tracker running at {url}")
    print(f"Private state: {args.state}")
    if not args.no_open:
        webbrowser.open(url)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
